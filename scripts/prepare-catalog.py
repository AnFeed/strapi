#!/usr/bin/env python3
"""Create a reviewable, draft-only catalog export from the supplied workbook.

Usage: python3 scripts/prepare-catalog.py /path/to/products.xlsx
Requires openpyxl. No Strapi connection or database write is made.
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


CATEGORY_CODES = {
    "Baromfitakarmányok": "poultry",
    "Sertéstakarmányok": "swine",
    "Szarvasmarhatakarmányok": "cattle",
    "Juhtakarmányok": "sheep",
    "Nyúltakarmányok": "rabbit",
    "Takarmányalapanyagok": "ingredients",
    "Vitamin- és ásványianyag-premixek": "premixes",
}


def split(value, *, shared=False):
    if value is None:
        return {"hu": None, "ro": None}
    value = str(value).strip()
    if " / " in value:
        hu, ro = value.split(" / ", 1)
        return {"hu": hu.strip(), "ro": ro.strip()}
    if shared:
        return {"hu": value, "ro": value}
    return {"hu": value, "ro": None}


def slug_pair(value):
    if value is None:
        return {"hu": None, "ro": None}
    value = str(value).strip()
    if " / " in value:
        return split(value)
    if "/ " in value:
        hu, ro = value.split("/ ", 1)
        return {"hu": hu.strip(), "ro": ro.strip()}
    # Brand names such as RoboAroMilk have the same slug in both languages.
    return split(value, shared=True)


def packages(value):
    text = str(value or "").lower().replace(" ", "")
    mapping = [("1kg", "kg_1"), ("5kg", "kg_5"),
               ("20kg", "kg_20"), ("25kg", "kg_25"),
               ("bigbag", "big_bag"), ("ömlesztett", "bulk")]
    return [{"size": code} for token, code in mapping if token in text]


def forms(value):
    text = str(value or "").lower()
    mapping = [("morzsolt", "crumbled"), ("lisztes", "meal"),
               ("granulált", "pellet")]
    return [{"form": code, "image": None} for token, code in mapping if token in text]


def code(value):
    return str(int(value)) if isinstance(value, (int, float)) else str(value).strip()


def prepare(workbook):
    categories = []
    seen = set()
    for row in workbook["Kategóriák és alkategóriák"].iter_rows(min_row=2, values_only=True):
        if not row[0] or row[0] in seen:
            continue
        seen.add(row[0])
        names = split(row[0])
        key = CATEGORY_CODES[names["hu"]]
        paths = split(row[2])
        categories.append({
            "categoryCode": key,
            "name": names,
            "slug": {lang: (paths[lang] or "").strip("/").split("/")[-1]
                     for lang in ("hu", "ro")},
            "introduction": split(row[4]),
            "seoTitle": split(row[5]),
            "metaDescription": split(row[6]),
            "publicationState": "draft",
        })

    products = []
    review = []
    for row_no, row in enumerate(workbook["Termékkártyák"].iter_rows(min_row=2, values_only=True), 2):
        if not row[2]:
            continue
        product_code = code(row[2])
        name = split(row[4], shared=row[4] in ("RoboAroMilk", "RoboProtMilk"))
        item = {
            "productCode": product_code,
            "categoryCode": CATEGORY_CODES[split(row[0])["hu"]],
            "subcategory": split(row[1], shared=str(row[1]).strip() == "-"),
            "availability": {"Üzlet": "in_store", "Előrendelés": "preorder"}.get(row[3]),
            "name": name,
            "slug": slug_pair(row[5]),
            "lifeStage": split(row[6]),
            "feedType": split(row[7]),
            "introduction": split(row[8]),
            "forms": forms(row[9]),
            "imageAlt": split(row[11], shared=row[11] in ("RoboAroMilk", "RoboProtMilk")),
            "packages": packages(row[12]),
            "ingredients": split(row[13]),
            "analyticalValues": str(row[14]).strip() if row[14] else None,
            "benefits": split(row[15]),
            "usage": split(row[16]),
            "labelFileName": str(row[17]).strip() if row[17] else None,
            "relatedProductCodes": [x.strip() for x in re.split(r",\s*", str(row[18])) if x.strip()]
                                   if row[18] else [],
            "seoTitle": split(row[19]),
            "metaDescription": split(row[20]),
            "publicationState": "draft",
        }
        if not item["slug"]["hu"] or not item["slug"]["ro"]:
            review.append({"row": row_no, "productCode": product_code, "issue": "missing localized slug"})
        if not item["ingredients"]["ro"]:
            review.append({"row": row_no, "productCode": product_code, "issue": "missing Romanian ingredients"})
        if not item["imageAlt"]["ro"]:
            review.append({"row": row_no, "productCode": product_code, "issue": "missing Romanian image alt text"})
        if not item["availability"]:
            review.append({"row": row_no, "productCode": product_code, "issue": "missing availability"})
        if not item["forms"]:
            review.append({"row": row_no, "productCode": product_code, "issue": "missing feed form"})
        if not item["analyticalValues"]:
            review.append({"row": row_no, "productCode": product_code, "issue": "missing analytical values"})
        if not item["labelFileName"]:
            review.append({"row": row_no, "productCode": product_code, "issue": "missing label PDF name"})
        products.append(item)

    assert len(categories) == 7 and len(products) == 52
    assert len({x["productCode"] for x in products}) == len(products)
    for locale in ("hu", "ro"):
        used = {}
        for item in products:
            slug = item["slug"][locale]
            if slug in used and slug is not None:
                review.append({"productCode": item["productCode"], "issue":
                               f"duplicate {locale} slug with {used[slug]}: {slug}"})
            used[slug] = item["productCode"]
    return {"source": "Product development workbook, 2026-10-14", "categories": categories,
            "products": products, "review": review}


if __name__ == "__main__":
    data = prepare(load_workbook(sys.argv[1], data_only=True))
    out = Path(__file__).resolve().parent.parent / "data" / "catalog-import-draft.json"
    out.parent.mkdir(exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"{len(data['categories'])} categories, {len(data['products'])} products, "
          f"{len(data['review'])} review items: {out}")
